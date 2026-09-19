import os
import sys
import pptx
from pptx.util import Inches as PInches, Pt as PPt
from pptx.dml.color import RGBColor as PRGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from reportlab.lib import pagesizes, colors
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas

SUBMISSION_DIR = r"D:\MIRROR\submission"
ASSETS_DIR = os.path.join(SUBMISSION_DIR, "assets")

IMG_OVERVIEW = os.path.join(ASSETS_DIR, "overview_flow_state.png")
IMG_TIMELINE = os.path.join(ASSETS_DIR, "activity_timeline.png")
IMG_PATTERNS = os.path.join(ASSETS_DIR, "behavioral_patterns_slm.png")
IMG_TRENDS = os.path.join(ASSETS_DIR, "trends_adaptive_baseline.png")

PPTX_PATH = os.path.join(SUBMISSION_DIR, "Mirror_Pitch_Deck.pptx")
PDF_PATH = os.path.join(SUBMISSION_DIR, "Mirror_Pitch_Deck.pdf")

# ==============================================================================
# 1. BUILD PPTX WITH REAL APPLICATION SCREENSHOTS & QUALCOMM THEME
# ==============================================================================
def build_winning_pptx():
    print("[1/2] Building Qualcomm Hackathon-Winning PPTX...")
    prs = pptx.Presentation()
    prs.slide_width = PInches(13.333)
    prs.slide_height = PInches(7.5)
    blank_layout = prs.slide_layouts[6]

    BG_COLOR = PRGBColor(15, 23, 42)        # Slate 900 #0F172A
    CARD_BG = PRGBColor(30, 41, 59)         # Slate 800 #1E293B
    BORDER_COLOR = PRGBColor(51, 65, 85)    # Slate 700 #334155
    ACCENT_CYAN = PRGBColor(14, 165, 233)   # Cyan #0EA5E9
    ACCENT_EMERALD = PRGBColor(16, 185, 129)# Emerald #10B981
    ACCENT_AMBER = PRGBColor(245, 158, 11)  # Amber #F59E0B
    ACCENT_PURPLE = PRGBColor(168, 85, 247) # Purple #A855F7
    TEXT_WHITE = PRGBColor(248, 250, 252)
    TEXT_MUTED = PRGBColor(148, 163, 184)

    def set_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_COLOR):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = PPt(1.5)
        return shape

    def add_header(slide, tag, title, tag_color=ACCENT_CYAN):
        tx = slide.shapes.add_textbox(PInches(0.8), PInches(0.5), PInches(11.7), PInches(1.1))
        tf = tx.text_frame
        p0 = tf.paragraphs[0]
        p0.text = tag.upper()
        p0.font.name = "Segoe UI"
        p0.font.size = PPt(11)
        p0.font.bold = True
        p0.font.color.rgb = tag_color
        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.name = "Segoe UI"
        p1.font.size = PPt(25)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE
        p1.space_before = PPt(2)

    # --------------------------------------------------------------------------
    # SLIDE 1: Title & Hackathon Hero
    # --------------------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)

    # Top badges
    b1 = add_card(s1, PInches(0.8), PInches(0.9), PInches(3.3), PInches(0.4), ACCENT_EMERALD, ACCENT_EMERALD)
    b1.text_frame.text = "QUALCOMM SNAPDRAGON X HACKATHON"
    p = b1.text_frame.paragraphs[0]
    p.font.name = "Segoe UI"; p.font.size = PPt(9.5); p.font.bold = True; p.font.color.rgb = BG_COLOR; p.alignment = PP_ALIGN.CENTER

    b2 = add_card(s1, PInches(4.3), PInches(0.9), PInches(3.1), PInches(0.4), CARD_BG, ACCENT_CYAN)
    b2.text_frame.text = "100% ON-DEVICE  |  ZERO CLOUD"
    p = b2.text_frame.paragraphs[0]
    p.font.name = "Segoe UI"; p.font.size = PPt(9.5); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN; p.alignment = PP_ALIGN.CENTER

    tx1 = s1.shapes.add_textbox(PInches(0.8), PInches(1.5), PInches(11.7), PInches(2.5))
    tf1 = tx1.text_frame; tf1.word_wrap = True
    p = tf1.paragraphs[0]; p.text = "MIRROR"; p.font.name = "Segoe UI"; p.font.size = PPt(54); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN
    p = tf1.add_paragraph(); p.text = "The Privacy-First, On-Device Digital Wellbeing Companion"; p.font.name = "Segoe UI"; p.font.size = PPt(26); p.font.bold = True; p.font.color.rgb = TEXT_WHITE
    p = tf1.add_paragraph(); p.text = "Unleashing the Qualcomm Hexagon NPU for Real-Time, Ethical Edge-AI Behavioral Intelligence on Windows 11 Copilot+ PCs"; p.font.name = "Segoe UI"; p.font.size = PPt(14.5); p.font.color.rgb = TEXT_MUTED; p.space_before = PPt(6)

    h_items = [
        ("SNAPDRAGON X NPU", "ONNX Runtime with Qualcomm QNN Execution Provider. 0.02ms P50 latency at <67 MB RAM.", ACCENT_CYAN),
        ("7 STANDOUT INNOVATIONS", "Adaptive Personal Baseline, Flow-State Deep Work, Local SLM RAG, Wellbeing Wallet, Audio Narrator.", ACCENT_EMERALD),
        ("MATHEMATICAL PRIVACY", "Zero cloud account, zero network sockets, zero keystrokes or URLs. 65/65 automated tests passing.", ACCENT_AMBER)
    ]
    for i, (head, desc, clr) in enumerate(h_items):
        c = add_card(s1, PInches(0.8 + i * 3.98), PInches(4.5), PInches(3.75), PInches(2.2))
        tf = c.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = head; p.font.name = "Segoe UI"; p.font.size = PPt(14); p.font.bold = True; p.font.color.rgb = clr
        p = tf.add_paragraph(); p.text = desc; p.font.name = "Segoe UI"; p.font.size = PPt(11); p.font.color.rgb = TEXT_MUTED; p.space_before = PPt(8)

    # --------------------------------------------------------------------------
    # SLIDE 2: The Problem: The Surveillance Dilemma
    # --------------------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_header(s2, "The Problem", "Screen-Time Trackers Shouldn't Turn Into Cloud Surveillance", PRGBColor(239, 68, 68))

    probs = [
        ("Invasive Telemetry", "Legacy tools log full window titles, URLs, keystrokes, and periodic screenshots—exposing passwords, banking, and private chats to remote servers.", "🚩"),
        ("Cloud Vulnerability", "Aggregating intimate behavioral habits in centralized cloud databases creates constant data leak risks and enables unconsented corporate profiling.", "☁️"),
        ("Guilt & Negative Shaming", "Conventional apps rely on intrusive modal timers and clinical medical labels ('burnout', 'addiction'), increasing anxiety rather than agency.", "⚠️")
    ]
    for i, (title, desc, icon) in enumerate(probs):
        c = add_card(s2, PInches(0.8 + i * 3.98), PInches(2.0), PInches(3.75), PInches(4.6))
        tf = c.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = icon; p.font.size = PPt(28)
        p = tf.add_paragraph(); p.text = title; p.font.name = "Segoe UI"; p.font.size = PPt(17); p.font.bold = True; p.font.color.rgb = TEXT_WHITE; p.space_before = PPt(10)
        p = tf.add_paragraph(); p.text = desc; p.font.name = "Segoe UI"; p.font.size = PPt(12); p.font.color.rgb = TEXT_MUTED; p.space_before = PPt(12)

    # --------------------------------------------------------------------------
    # SLIDE 3: The Mirror Solution & Architecture
    # --------------------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_header(s3, "Architecture & Pipeline", "100% On-Device: From Win32 Hooks to Snapdragon X NPU", ACCENT_CYAN)

    steps = [
        ("1. Coarse Win32 Hook", "SetWinEventHook tracks foreground transitions; GetLastInputInfo tracks idle time. Zero window titles or inputs read."),
        ("2. Temporal Matrix", "Encodes sequences into [60×10] normalized feature matrices capturing duration and transition velocity."),
        ("3. Qualcomm QNN NPU", "ONNX Runtime executes 1D-CNN temporal sequence model on Snapdragon X Hexagon NPU in 0.02ms P50 latency."),
        ("4. Signal Fusion Engine", "Blends deterministic rules with local neural probabilities to guarantee 0 false positives without cloud checks."),
        ("5. SQLite in WAL Mode", "Zero-cloud persistence with transactional integrity, fast local queries, and 0-residue data purge."),
        ("6. Windows 11 Fluent UI", "WinUI 3 native desktop app featuring custom Mica backdrop, real-time live cards, and dark theme.")
    ]
    for i, (title, desc) in enumerate(steps):
        row = i // 3; col = i % 3
        c = add_card(s3, PInches(0.8 + col * 3.98), PInches(1.9 + row * 2.5), PInches(3.75), PInches(2.2))
        tf = c.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = title; p.font.name = "Segoe UI"; p.font.size = PPt(14); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN
        p = tf.add_paragraph(); p.text = desc; p.font.name = "Segoe UI"; p.font.size = PPt(10.5); p.font.color.rgb = TEXT_MUTED; p.space_before = PPt(6)

    # --------------------------------------------------------------------------
    # SLIDE 4: Live UI Showcase — Flow State & Live Tracking (SCREENSHOT 1)
    # --------------------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    add_header(s4, "Live UI Showcase — Positive Reinforcement", "Flow-State Recognition & Real-Time Tracking", ACCENT_EMERALD)

    # Left: Real App Screenshot with rounded border card
    img_card = add_card(s4, PInches(0.8), PInches(1.8), PInches(7.8), PInches(4.9), CARD_BG, ACCENT_EMERALD)
    s4.shapes.add_picture(IMG_OVERVIEW, PInches(0.9), PInches(1.9), PInches(7.6), PInches(4.7))

    # Right: Callout highlights
    callouts_s4 = [
        ("Flow State Achieved Card", "Detects ≥60 min uninterrupted deep work with 0 distractions. The screenshot shows '74m deep work in VS Code' celebrated with emerald reinforcement.", ACCENT_EMERALD),
        ("Real-Time Focus Card", "Displays live foreground app with an active focus duration ticking every moment ('02:06:24') and second-level precision.", ACCENT_CYAN),
        ("Voice-Narrated Summary", "Native Windows offline speech synthesis ('Listen to Today's Summary' button) delivers calm, human daily audio briefs.", ACCENT_PURPLE)
    ]
    for i, (h, d, clr) in enumerate(callouts_s4):
        c = add_card(s4, PInches(8.8), PInches(1.8 + i * 1.68), PInches(3.7), PInches(1.55))
        tf = c.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = h; p.font.name = "Segoe UI"; p.font.size = PPt(12.5); p.font.bold = True; p.font.color.rgb = clr
        p = tf.add_paragraph(); p.text = d; p.font.name = "Segoe UI"; p.font.size = PPt(9.5); p.font.color.rgb = TEXT_MUTED; p.space_before = PPt(4)

    # --------------------------------------------------------------------------
    # SLIDE 5: Live UI Showcase — Local SLM RAG & Active Recalibration (SCREENSHOT 2)
    # --------------------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_header(s5, "Live UI Showcase — Edge Generative AI", "Explain with Local AI (SLM / RAG) & Active Recalibration", ACCENT_CYAN)

    img_card5 = add_card(s5, PInches(0.8), PInches(1.8), PInches(7.8), PInches(4.9), CARD_BG, ACCENT_CYAN)
    s5.shapes.add_picture(IMG_PATTERNS, PInches(0.9), PInches(1.9), PInches(7.6), PInches(4.7))

    callouts_s5 = [
        ("On-Device SLM RAG Reasoning", "Extracts a 4-hour local SQLite context window around flagged patterns. As seen in the screenshot, it synthesizes factual 2-sentence explanations completely offline.", ACCENT_CYAN),
        ("Active Recalibration Loop", "Users can click 'This isn't accurate' to instantly bump sensitivity thresholds by +15%, teaching Mirror their habits without cloud retraining.", ACCENT_AMBER),
        ("Non-Clinical Objective Framing", "Enforces strict ethical guidelines prohibiting diagnostic labels ('burnout', 'addiction'), presenting purely observational evidence.", ACCENT_EMERALD)
    ]
    for i, (h, d, clr) in enumerate(callouts_s5):
        c = add_card(s5, PInches(8.8), PInches(1.8 + i * 1.68), PInches(3.7), PInches(1.55))
        tf = c.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = h; p.font.name = "Segoe UI"; p.font.size = PPt(12.5); p.font.bold = True; p.font.color.rgb = clr
        p = tf.add_paragraph(); p.text = d; p.font.name = "Segoe UI"; p.font.size = PPt(9.5); p.font.color.rgb = TEXT_MUTED; p.space_before = PPt(4)

    # --------------------------------------------------------------------------
    # SLIDE 6: Live UI Showcase — Adaptive Statistical Baseline (SCREENSHOT 3)
    # --------------------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)
    add_header(s6, "Live UI Showcase — Personalization Without Profiling", "Adaptive Personal Baseline & Statistical Envelope", ACCENT_EMERALD)

    img_card6 = add_card(s6, PInches(0.8), PInches(1.8), PInches(7.8), PInches(4.9), CARD_BG, ACCENT_EMERALD)
    s6.shapes.add_picture(IMG_TRENDS, PInches(0.9), PInches(1.9), PInches(7.6), PInches(4.7))

    callouts_s6 = [
        ("Rolling Statistical Envelope", "Replaces rigid arbitrary limits with a 14-day rolling median (6.5h) and natural variance standard deviation (±1.1h) derived entirely on-device.", ACCENT_EMERALD),
        ("Strict >2σ Flagging Gate", "Anomalous patterns trigger only if usage exceeds two standard deviations (>2σ = 8.7h), eliminating false alerts during typical busy days.", ACCENT_CYAN),
        ("Day 7 Calibration Phase Gate", "Strictly gates pattern alerts during the first 7 days of installation while Mirror calmly learns baseline habits.", PRGBColor(236, 72, 153))
    ]
    for i, (h, d, clr) in enumerate(callouts_s6):
        c = add_card(s6, PInches(8.8), PInches(1.8 + i * 1.68), PInches(3.7), PInches(1.55))
        tf = c.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = h; p.font.name = "Segoe UI"; p.font.size = PPt(12.5); p.font.bold = True; p.font.color.rgb = clr
        p = tf.add_paragraph(); p.text = d; p.font.name = "Segoe UI"; p.font.size = PPt(9.5); p.font.color.rgb = TEXT_MUTED; p.space_before = PPt(4)

    # --------------------------------------------------------------------------
    # SLIDE 7: Live UI Showcase — Real-Time Activity Timeline (SCREENSHOT 4)
    # --------------------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7)
    add_header(s7, "Live UI Showcase — Sub-Second Observability", "Chronological Activity Timeline Updated Every Moment", ACCENT_PURPLE)

    img_card7 = add_card(s7, PInches(0.8), PInches(1.8), PInches(7.8), PInches(4.9), CARD_BG, ACCENT_PURPLE)
    s7.shapes.add_picture(IMG_TIMELINE, PInches(0.9), PInches(1.9), PInches(7.6), PInches(4.7))

    callouts_s7 = [
        ("Updated Every Moment", "Dynamically prepends the currently focused app as 'LIVE NOW' with active duration, smoothly archiving completed sessions on window switch.", ACCENT_PURPLE),
        ("Categorical Classification", "Maps processes to normalized categories (Development, Communication, Browsing) using an embedded offline heuristic identity resolver.", ACCENT_CYAN),
        ("Zero Private Content Ingestion", "Only logs (Timestamp, AppKey, Process, Category, Duration, CloseReason). Zero document titles, URLs, or text inputs ever recorded.", ACCENT_EMERALD)
    ]
    for i, (h, d, clr) in enumerate(callouts_s7):
        c = add_card(s7, PInches(8.8), PInches(1.8 + i * 1.68), PInches(3.7), PInches(1.55))
        tf = c.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = h; p.font.name = "Segoe UI"; p.font.size = PPt(12.5); p.font.bold = True; p.font.color.rgb = clr
        p = tf.add_paragraph(); p.text = d; p.font.name = "Segoe UI"; p.font.size = PPt(9.5); p.font.color.rgb = TEXT_MUTED; p.space_before = PPt(4)

    # --------------------------------------------------------------------------
    # SLIDE 8: Why Snapdragon X NPU Wins (Benchmarks)
    # --------------------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8)
    add_header(s8, "Snapdragon X NPU Hardware Acceleration", "Sub-Millisecond On-Device Inference Benchmarks", ACCENT_CYAN)

    bench_cards = [
        ("0.02 ms", "STEADY-STATE P50 LATENCY", "Snapdragon X Hexagon NPU executes model evaluation in 20 microseconds.", ACCENT_CYAN),
        ("0.37 ms", "WARMUP COLD-START", "Instantaneous initialization on session launch with Qualcomm QNN EP.", ACCENT_EMERALD),
        ("< 67 MB", "ACTIVE RAM FOOTPRINT", "Ultra-lightweight process footprint designed for perpetual background operation.", ACCENT_AMBER),
        ("< 0.1%", "BACKGROUND CPU USAGE", "Virtually zero battery penalty on Snapdragon X Copilot+ laptops.", ACCENT_PURPLE)
    ]
    for i, (metric, label, desc, clr) in enumerate(bench_cards):
        c = add_card(s8, PInches(0.8 + i * 2.98), PInches(2.0), PInches(2.78), PInches(2.2))
        tf = c.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = metric; p.font.name = "Segoe UI"; p.font.size = PPt(32); p.font.bold = True; p.font.color.rgb = clr
        p = tf.add_paragraph(); p.text = label; p.font.name = "Segoe UI"; p.font.size = PPt(9.5); p.font.bold = True; p.font.color.rgb = TEXT_WHITE; p.space_before = PPt(2)
        p = tf.add_paragraph(); p.text = desc; p.font.name = "Segoe UI"; p.font.size = PPt(10); p.font.color.rgb = TEXT_MUTED; p.space_before = PPt(6)

    # Lower comparison table card
    c_tbl = add_card(s8, PInches(0.8), PInches(4.5), PInches(11.7), PInches(2.3))
    tf = c_tbl.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "THREE-TIER HARDWARE EXECUTION PROVIDER FALLBACK"; p.font.name = "Segoe UI"; p.font.size = PPt(12); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN
    p = tf.add_paragraph(); p.text = "1. Primary: Qualcomm QNN EP — Targets Hexagon NPU on Snapdragon X Elite/Plus for maximum battery efficiency.\n2. Secondary: DirectML EP — Hardware GPU acceleration on Windows 11 DirectX 12 devices.\n3. Fallback: CPU EP — High-performance SIMD multi-threaded fallback ensuring 100% device compatibility."; p.font.name = "Segoe UI"; p.font.size = PPt(11.5); p.font.color.rgb = TEXT_WHITE; p.space_before = PPt(6)

    # --------------------------------------------------------------------------
    # SLIDE 9: Security, Privacy & Verifiable Audit
    # --------------------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    set_bg(s9)
    add_header(s9, "Verification & Data Sovereignty", "Mathematical Proof: Zero Network & Encrypted Wallet", ACCENT_EMERALD)

    audits = [
        ("65 / 65 Automated Tests Passing", "100% test pass rate across 6 test suites: Core, Tracking, Persistence, Analytics, Inference, and Privacy with 0 failures.", "✅", ACCENT_EMERALD),
        ("Reflection Privacy Guard", "Automated code inspection dynamically verifies zero references to System.Net, HttpClient, or sockets across all assemblies.", "🛡️", ACCENT_CYAN),
        ("AES-256-GCM 'Wellbeing Wallet'", "Military-grade password-protected archive (.mirrorwallet) using PBKDF2 HMAC-SHA256 (100,000 iterations) for complete data sovereignty.", "🔐", ACCENT_PURPLE),
        ("Offline QuestPDF Report Card", "Generates professional executive vector PDF reports directly to Documents\\Mirror with zero external service calls.", "📄", ACCENT_AMBER)
    ]
    for i, (title, desc, icon, clr) in enumerate(audits):
        row = i // 2; col = i % 2
        c = add_card(s9, PInches(0.8 + col * 5.95), PInches(2.0 + row * 2.45), PInches(5.75), PInches(2.15))
        tf = c.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = f"{icon}  {title}"; p.font.name = "Segoe UI"; p.font.size = PPt(15); p.font.bold = True; p.font.color.rgb = clr
        p = tf.add_paragraph(); p.text = desc; p.font.name = "Segoe UI"; p.font.size = PPt(11.5); p.font.color.rgb = TEXT_MUTED; p.space_before = PPt(8)

    # --------------------------------------------------------------------------
    # SLIDE 10: Why Mirror Wins the Qualcomm Hackathon
    # --------------------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    set_bg(s10)

    c_win = add_card(s10, PInches(0.8), PInches(0.8), PInches(11.7), PInches(5.9), CARD_BG, ACCENT_CYAN)
    tf = c_win.text_frame; tf.word_wrap = True

    p = tf.paragraphs[0]; p.text = "THE QUALCOMM HACKATHON WINNING PROPOSITION"; p.font.name = "Segoe UI"; p.font.size = PPt(12); p.font.bold = True; p.font.color.rgb = ACCENT_CYAN
    p = tf.add_paragraph(); p.text = "Why Mirror is the Flagship Showcase for Snapdragon X Copilot+ PCs"; p.font.name = "Segoe UI"; p.font.size = PPt(26); p.font.bold = True; p.font.color.rgb = TEXT_WHITE; p.space_before = PPt(4)

    reasons = [
        ("1. Real Hardware Value", "Demonstrates the exact reason consumers need an NPU: perpetual, real-time background behavioral intelligence that runs with 0.02ms latency and zero battery impact."),
        ("2. Production-Grade Execution", "Not a prototype or concept. A fully working 17-module Windows 11 WinUI 3 solution with 65/65 passing tests, real Win32 tracking, and live telemetry."),
        ("3. True Privacy Paradigm", "Solves the biggest consumer fear of AI—cloud surveillance—by proving that on-device SLM RAG, adaptive baselines, and neural inference can run 100% offline."),
        ("4. Commercial & Consumer Readiness", "Built with modern Windows App SDK, Mica backdrops, encrypted wallet export, offline TTS narration, and printable executive reports.")
    ]
    for h, d in reasons:
        p = tf.add_paragraph()
        p.text = f"{h}: "
        p.font.name = "Segoe UI"; p.font.size = PPt(12); p.font.bold = True; p.font.color.rgb = ACCENT_EMERALD; p.space_before = PPt(12)
        r2 = p.add_run()
        r2.text = d
        r2.font.name = "Segoe UI"; r2.font.size = PPt(12); r2.font.color.rgb = TEXT_MUTED

    prs.save(PPTX_PATH)
    print(f"Saved 10-slide winning PPTX with real screenshots to {PPTX_PATH}")

# ==============================================================================
# 2. BUILD HIGH-FIDELITY PDF PRESENTATION (Matching 16:9 Deck)
# ==============================================================================
def build_winning_pdf():
    print("[2/2] Generating Qualcomm Hackathon-Winning 16:9 PDF...")
    width, height = 13.333 * inch, 7.5 * inch
    c = canvas.Canvas(PDF_PATH, pagesize=(width, height))

    BG_SLATE_900 = colors.HexColor('#0F172A')
    CARD_SLATE_800 = colors.HexColor('#1E293B')
    BORDER_SLATE_700 = colors.HexColor('#334155')
    ACCENT_CYAN = colors.HexColor('#0EA5E9')
    ACCENT_EMERALD = colors.HexColor('#10B981')
    ACCENT_AMBER = colors.HexColor('#F59E0B')
    ACCENT_PURPLE = colors.HexColor('#A855F7')
    ACCENT_RED = colors.HexColor('#EF4444')
    TEXT_WHITE = colors.HexColor('#F8FAFC')
    TEXT_MUTED = colors.HexColor('#94A3B8')

    def draw_bg():
        c.setFillColor(BG_SLATE_900)
        c.rect(0, 0, width, height, fill=1, stroke=0)

    def draw_card(x, y, w, h, bg=CARD_SLATE_800, border=BORDER_SLATE_700, r=8):
        c.setFillColor(bg)
        c.setStrokeColor(border)
        c.setLineWidth(1.5)
        c.roundRect(x, y, w, h, r, fill=1, stroke=1)

    def draw_header(tag, title, tag_clr=ACCENT_CYAN):
        c.setFont("Helvetica-Bold", 11)
        c.setFillColor(tag_clr)
        c.drawString(0.8 * inch, height - 0.8 * inch, tag.upper())
        c.setFont("Helvetica-Bold", 24)
        c.setFillColor(TEXT_WHITE)
        c.drawString(0.8 * inch, height - 1.25 * inch, title)

    # -------------------------------------------------------------
    # SLIDE 1: Title
    # -------------------------------------------------------------
    draw_bg()
    draw_card(0.8 * inch, height - 1.35 * inch, 3.3 * inch, 0.4 * inch, bg=ACCENT_EMERALD, border=ACCENT_EMERALD, r=6)
    c.setFont("Helvetica-Bold", 9.5); c.setFillColor(BG_SLATE_900)
    c.drawCentredString(0.8 * inch + 1.65 * inch, height - 1.22 * inch, "QUALCOMM SNAPDRAGON X HACKATHON")

    draw_card(4.3 * inch, height - 1.35 * inch, 3.1 * inch, 0.4 * inch, bg=CARD_SLATE_800, border=ACCENT_CYAN, r=6)
    c.setFont("Helvetica-Bold", 9.5); c.setFillColor(ACCENT_CYAN)
    c.drawCentredString(4.3 * inch + 1.55 * inch, height - 1.22 * inch, "100% ON-DEVICE  |  ZERO CLOUD")

    c.setFont("Helvetica-Bold", 54); c.setFillColor(ACCENT_CYAN)
    c.drawString(0.8 * inch, height - 2.4 * inch, "MIRROR")

    c.setFont("Helvetica-Bold", 25); c.setFillColor(TEXT_WHITE)
    c.drawString(0.8 * inch, height - 2.95 * inch, "The Privacy-First, On-Device Digital Wellbeing Companion")

    c.setFont("Helvetica", 14); c.setFillColor(TEXT_MUTED)
    c.drawString(0.8 * inch, height - 3.4 * inch, "Unleashing Qualcomm Hexagon NPU for Real-Time, Ethical Edge-AI on Windows 11 Copilot+ PCs")

    h_items = [
        ("SNAPDRAGON X NPU", "ONNX Runtime with Qualcomm QNN Execution Provider. 0.02ms P50 latency at <67 MB RAM.", ACCENT_CYAN),
        ("7 STANDOUT INNOVATIONS", "Adaptive Baseline, Flow-State Deep Work, Local SLM RAG, Wellbeing Wallet, Audio Narrator.", ACCENT_EMERALD),
        ("MATHEMATICAL PRIVACY", "Zero cloud account, zero network sockets, zero keystrokes or URLs. 65/65 automated tests passing.", ACCENT_AMBER)
    ]
    for i, (head, desc, clr) in enumerate(h_items):
        x = 0.8 * inch + i * 3.98 * inch
        y = 0.8 * inch
        draw_card(x, y, 3.75 * inch, 2.2 * inch)
        c.setFont("Helvetica-Bold", 14); c.setFillColor(clr); c.drawString(x + 18, y + 2.2 * inch - 30, head)
        c.setFont("Helvetica", 11); c.setFillColor(TEXT_MUTED)
        words = desc.split(); lines = []; cur = []
        for w in words:
            if len(" ".join(cur + [w])) < 35: cur.append(w)
            else: lines.append(" ".join(cur)); cur = [w]
        if cur: lines.append(" ".join(cur))
        for li, line in enumerate(lines):
            c.drawString(x + 18, y + 2.2 * inch - 60 - li * 18, line)
    c.showPage()

    # -------------------------------------------------------------
    # SLIDE 2: Problem
    # -------------------------------------------------------------
    draw_bg()
    draw_header("The Problem", "Screen-Time Trackers Shouldn't Turn Into Cloud Surveillance", ACCENT_RED)
    probs = [
        ("Invasive Telemetry", "Legacy tools log full window titles, URLs, keystrokes, and periodic screenshots—exposing passwords, banking, and private chats to remote servers.", ACCENT_RED),
        ("Cloud Vulnerability", "Aggregating intimate behavioral habits in centralized cloud databases creates constant data leak risks and enables unconsented corporate profiling.", ACCENT_AMBER),
        ("Guilt & Negative Shaming", "Conventional apps rely on intrusive modal timers and clinical medical labels ('burnout', 'addiction'), increasing anxiety rather than agency.", colors.HexColor('#F87171'))
    ]
    for i, (title, desc, clr) in enumerate(probs):
        x = 0.8 * inch + i * 3.98 * inch
        y = 0.9 * inch
        draw_card(x, y, 3.75 * inch, 4.6 * inch)
        c.setFont("Helvetica-Bold", 18); c.setFillColor(clr); c.drawString(x + 20, y + 4.6 * inch - 40, f"0{i+1}")
        c.setFont("Helvetica-Bold", 16); c.setFillColor(TEXT_WHITE); c.drawString(x + 20, y + 4.6 * inch - 75, title)
        c.setFont("Helvetica", 11.5); c.setFillColor(TEXT_MUTED)
        words = desc.split(); lines = []; cur = []
        for w in words:
            if len(" ".join(cur + [w])) < 32: cur.append(w)
            else: lines.append(" ".join(cur)); cur = [w]
        if cur: lines.append(" ".join(cur))
        for li, line in enumerate(lines):
            c.drawString(x + 20, y + 4.6 * inch - 110 - li * 20, line)
    c.showPage()

    # -------------------------------------------------------------
    # SLIDE 3: Architecture
    # -------------------------------------------------------------
    draw_bg()
    draw_header("Architecture & Pipeline", "100% On-Device: From Win32 Hooks to Snapdragon X NPU", ACCENT_CYAN)
    steps = [
        ("1. Coarse Win32 Hook", "SetWinEventHook tracks foreground transitions; GetLastInputInfo tracks idle time. Zero window titles or inputs read."),
        ("2. Temporal Matrix", "Encodes sequences into [60×10] normalized feature matrices capturing duration and transition velocity."),
        ("3. Qualcomm QNN NPU", "ONNX Runtime executes 1D-CNN temporal sequence model on Snapdragon X Hexagon NPU in 0.02ms P50 latency."),
        ("4. Signal Fusion Engine", "Blends deterministic rules with local neural probabilities to guarantee 0 false positives without cloud checks."),
        ("5. SQLite in WAL Mode", "Zero-cloud persistence with transactional integrity, fast local queries, and 0-residue data purge."),
        ("6. Windows 11 Fluent UI", "WinUI 3 native desktop app featuring custom Mica backdrop, real-time live cards, and dark theme.")
    ]
    for i, (title, desc) in enumerate(steps):
        row = i // 3; col = i % 3
        x = 0.8 * inch + col * 3.98 * inch
        y = height - 1.8 * inch - (row + 1) * 2.5 * inch
        draw_card(x, y, 3.75 * inch, 2.2 * inch)
        c.setFont("Helvetica-Bold", 14); c.setFillColor(ACCENT_CYAN); c.drawString(x + 18, y + 2.2 * inch - 30, title)
        c.setFont("Helvetica", 10.5); c.setFillColor(TEXT_MUTED)
        words = desc.split(); lines = []; cur = []
        for w in words:
            if len(" ".join(cur + [w])) < 35: cur.append(w)
            else: lines.append(" ".join(cur)); cur = [w]
        if cur: lines.append(" ".join(cur))
        for li, line in enumerate(lines):
            c.drawString(x + 18, y + 2.2 * inch - 56 - li * 17, line)
    c.showPage()

    # -------------------------------------------------------------
    # SLIDE 4: Live UI — Flow State (IMG 1)
    # -------------------------------------------------------------
    draw_bg()
    draw_header("Live UI Showcase — Positive Reinforcement", "Flow-State Recognition & Real-Time Tracking", ACCENT_EMERALD)
    draw_card(0.8 * inch, 0.8 * inch, 7.8 * inch, 4.9 * inch, border=ACCENT_EMERALD)
    c.drawImage(IMG_OVERVIEW, 0.9 * inch, 0.9 * inch, 7.6 * inch, 4.7 * inch, preserveAspectRatio=True)

    callouts_s4 = [
        ("Flow State Achieved Card", "Detects ≥60 min uninterrupted deep work with 0 distractions. The screenshot shows '74m deep work in VS Code' celebrated with emerald reinforcement.", ACCENT_EMERALD),
        ("Real-Time Focus Card", "Displays live foreground app with an active focus duration ticking every moment ('02:06:24') and second-level precision.", ACCENT_CYAN),
        ("Voice-Narrated Summary", "Native Windows offline speech synthesis ('Listen to Today's Summary' button) delivers calm, human daily audio briefs.", ACCENT_PURPLE)
    ]
    for i, (h, d, clr) in enumerate(callouts_s4):
        x = 8.8 * inch; y = height - 1.7 * inch - (i + 1) * 1.68 * inch
        draw_card(x, y, 3.7 * inch, 1.55 * inch)
        c.setFont("Helvetica-Bold", 12.5); c.setFillColor(clr); c.drawString(x + 14, y + 1.55 * inch - 26, h)
        c.setFont("Helvetica", 9.5); c.setFillColor(TEXT_MUTED)
        words = d.split(); lines = []; cur = []
        for w in words:
            if len(" ".join(cur + [w])) < 38: cur.append(w)
            else: lines.append(" ".join(cur)); cur = [w]
        if cur: lines.append(" ".join(cur))
        for li, line in enumerate(lines):
            c.drawString(x + 14, y + 1.55 * inch - 50 - li * 15, line)
    c.showPage()

    # -------------------------------------------------------------
    # SLIDE 5: Live UI — Local SLM RAG (IMG 2)
    # -------------------------------------------------------------
    draw_bg()
    draw_header("Live UI Showcase — Edge Generative AI", "Explain with Local AI (SLM / RAG) & Active Recalibration", ACCENT_CYAN)
    draw_card(0.8 * inch, 0.8 * inch, 7.8 * inch, 4.9 * inch, border=ACCENT_CYAN)
    c.drawImage(IMG_PATTERNS, 0.9 * inch, 0.9 * inch, 7.6 * inch, 4.7 * inch, preserveAspectRatio=True)

    callouts_s5 = [
        ("On-Device SLM RAG Reasoning", "Extracts a 4-hour local SQLite context window around flagged patterns. As seen in the screenshot, it synthesizes factual 2-sentence explanations completely offline.", ACCENT_CYAN),
        ("Active Recalibration Loop", "Users can click 'This isn't accurate' to instantly bump sensitivity thresholds by +15%, teaching Mirror their habits without cloud retraining.", ACCENT_AMBER),
        ("Non-Clinical Objective Framing", "Enforces strict ethical guidelines prohibiting diagnostic labels ('burnout', 'addiction'), presenting purely observational evidence.", ACCENT_EMERALD)
    ]
    for i, (h, d, clr) in enumerate(callouts_s5):
        x = 8.8 * inch; y = height - 1.7 * inch - (i + 1) * 1.68 * inch
        draw_card(x, y, 3.7 * inch, 1.55 * inch)
        c.setFont("Helvetica-Bold", 12.5); c.setFillColor(clr); c.drawString(x + 14, y + 1.55 * inch - 26, h)
        c.setFont("Helvetica", 9.5); c.setFillColor(TEXT_MUTED)
        words = d.split(); lines = []; cur = []
        for w in words:
            if len(" ".join(cur + [w])) < 38: cur.append(w)
            else: lines.append(" ".join(cur)); cur = [w]
        if cur: lines.append(" ".join(cur))
        for li, line in enumerate(lines):
            c.drawString(x + 14, y + 1.55 * inch - 50 - li * 15, line)
    c.showPage()

    # -------------------------------------------------------------
    # SLIDE 6: Live UI — Adaptive Baseline (IMG 3)
    # -------------------------------------------------------------
    draw_bg()
    draw_header("Live UI Showcase — Personalization Without Profiling", "Adaptive Personal Baseline & Statistical Envelope", ACCENT_EMERALD)
    draw_card(0.8 * inch, 0.8 * inch, 7.8 * inch, 4.9 * inch, border=ACCENT_EMERALD)
    c.drawImage(IMG_TRENDS, 0.9 * inch, 0.9 * inch, 7.6 * inch, 4.7 * inch, preserveAspectRatio=True)

    callouts_s6 = [
        ("Rolling Statistical Envelope", "Replaces rigid arbitrary limits with a 14-day rolling median (6.5h) and natural variance standard deviation (±1.1h) derived entirely on-device.", ACCENT_EMERALD),
        ("Strict >2σ Flagging Gate", "Anomalous patterns trigger only if usage exceeds two standard deviations (>2σ = 8.7h), eliminating false alerts during typical busy days.", ACCENT_CYAN),
        ("Day 7 Calibration Phase Gate", "Strictly gates pattern alerts during the first 7 days of installation while Mirror calmly learns baseline habits.", colors.HexColor('#EC4899'))
    ]
    for i, (h, d, clr) in enumerate(callouts_s6):
        x = 8.8 * inch; y = height - 1.7 * inch - (i + 1) * 1.68 * inch
        draw_card(x, y, 3.7 * inch, 1.55 * inch)
        c.setFont("Helvetica-Bold", 12.5); c.setFillColor(clr); c.drawString(x + 14, y + 1.55 * inch - 26, h)
        c.setFont("Helvetica", 9.5); c.setFillColor(TEXT_MUTED)
        words = d.split(); lines = []; cur = []
        for w in words:
            if len(" ".join(cur + [w])) < 38: cur.append(w)
            else: lines.append(" ".join(cur)); cur = [w]
        if cur: lines.append(" ".join(cur))
        for li, line in enumerate(lines):
            c.drawString(x + 14, y + 1.55 * inch - 50 - li * 15, line)
    c.showPage()

    # -------------------------------------------------------------
    # SLIDE 7: Live UI — Timeline (IMG 4)
    # -------------------------------------------------------------
    draw_bg()
    draw_header("Live UI Showcase — Sub-Second Observability", "Chronological Activity Timeline Updated Every Moment", ACCENT_PURPLE)
    draw_card(0.8 * inch, 0.8 * inch, 7.8 * inch, 4.9 * inch, border=ACCENT_PURPLE)
    c.drawImage(IMG_TIMELINE, 0.9 * inch, 0.9 * inch, 7.6 * inch, 4.7 * inch, preserveAspectRatio=True)

    callouts_s7 = [
        ("Updated Every Moment", "Dynamically prepends the currently focused app as 'LIVE NOW' with active duration, smoothly archiving completed sessions on window switch.", ACCENT_PURPLE),
        ("Categorical Classification", "Maps processes to normalized categories (Development, Communication, Browsing) using an embedded offline heuristic identity resolver.", ACCENT_CYAN),
        ("Zero Private Content Ingestion", "Only logs (Timestamp, AppKey, Process, Category, Duration, CloseReason). Zero document titles, URLs, or text inputs ever recorded.", ACCENT_EMERALD)
    ]
    for i, (h, d, clr) in enumerate(callouts_s7):
        x = 8.8 * inch; y = height - 1.7 * inch - (i + 1) * 1.68 * inch
        draw_card(x, y, 3.7 * inch, 1.55 * inch)
        c.setFont("Helvetica-Bold", 12.5); c.setFillColor(clr); c.drawString(x + 14, y + 1.55 * inch - 26, h)
        c.setFont("Helvetica", 9.5); c.setFillColor(TEXT_MUTED)
        words = d.split(); lines = []; cur = []
        for w in words:
            if len(" ".join(cur + [w])) < 38: cur.append(w)
            else: lines.append(" ".join(cur)); cur = [w]
        if cur: lines.append(" ".join(cur))
        for li, line in enumerate(lines):
            c.drawString(x + 14, y + 1.55 * inch - 50 - li * 15, line)
    c.showPage()

    # -------------------------------------------------------------
    # SLIDE 8: Benchmarks
    # -------------------------------------------------------------
    draw_bg()
    draw_header("Snapdragon X NPU Hardware Acceleration", "Sub-Millisecond On-Device Inference Benchmarks", ACCENT_CYAN)

    bench_cards = [
        ("0.02 ms", "STEADY-STATE P50 LATENCY", "Snapdragon X Hexagon NPU executes model evaluation in 20 microseconds.", ACCENT_CYAN),
        ("0.37 ms", "WARMUP COLD-START", "Instantaneous initialization on session launch with Qualcomm QNN EP.", ACCENT_EMERALD),
        ("< 67 MB", "ACTIVE RAM FOOTPRINT", "Ultra-lightweight process footprint designed for perpetual background operation.", ACCENT_AMBER),
        ("< 0.1%", "BACKGROUND CPU USAGE", "Virtually zero battery penalty on Snapdragon X Copilot+ laptops.", ACCENT_PURPLE)
    ]
    for i, (metric, label, desc, clr) in enumerate(bench_cards):
        x = 0.8 * inch + i * 2.98 * inch
        y = height - 2.0 * inch - 2.2 * inch
        draw_card(x, y, 2.78 * inch, 2.2 * inch)
        c.setFont("Helvetica-Bold", 32); c.setFillColor(clr); c.drawString(x + 16, y + 2.2 * inch - 46, metric)
        c.setFont("Helvetica-Bold", 9.5); c.setFillColor(TEXT_WHITE); c.drawString(x + 16, y + 2.2 * inch - 74, label)
        c.setFont("Helvetica", 10); c.setFillColor(TEXT_MUTED)
        words = desc.split(); lines = []; cur = []
        for w in words:
            if len(" ".join(cur + [w])) < 26: cur.append(w)
            else: lines.append(" ".join(cur)); cur = [w]
        if cur: lines.append(" ".join(cur))
        for li, line in enumerate(lines):
            c.drawString(x + 16, y + 2.2 * inch - 110 - li * 16, line)

    y_tbl = 0.8 * inch
    draw_card(0.8 * inch, y_tbl, 11.7 * inch, 2.2 * inch)
    c.setFont("Helvetica-Bold", 12); c.setFillColor(ACCENT_CYAN); c.drawString(1.1 * inch, y_tbl + 2.2 * inch - 28, "THREE-TIER HARDWARE EXECUTION PROVIDER FALLBACK")
    c.setFont("Helvetica", 11.5); c.setFillColor(TEXT_WHITE)
    c.drawString(1.1 * inch, y_tbl + 2.2 * inch - 56, "1. Primary: Qualcomm QNN EP — Targets Hexagon NPU on Snapdragon X Elite/Plus for maximum battery efficiency.")
    c.drawString(1.1 * inch, y_tbl + 2.2 * inch - 82, "2. Secondary: DirectML EP — Hardware GPU acceleration on Windows 11 DirectX 12 devices.")
    c.drawString(1.1 * inch, y_tbl + 2.2 * inch - 108, "3. Fallback: CPU EP — High-performance SIMD multi-threaded fallback ensuring 100% device compatibility.")
    c.showPage()

    # -------------------------------------------------------------
    # SLIDE 9: Security
    # -------------------------------------------------------------
    draw_bg()
    draw_header("Verification & Data Sovereignty", "Mathematical Proof: Zero Network & Encrypted Wallet", ACCENT_EMERALD)
    audits = [
        ("65 / 65 Automated Tests Passing", "100% test pass rate across 6 test suites: Core, Tracking, Persistence, Analytics, Inference, and Privacy with 0 failures.", ACCENT_EMERALD),
        ("Reflection Privacy Guard", "Automated code inspection dynamically verifies zero references to System.Net, HttpClient, or sockets across all assemblies.", ACCENT_CYAN),
        ("AES-256-GCM 'Wellbeing Wallet'", "Military-grade password-protected archive (.mirrorwallet) using PBKDF2 HMAC-SHA256 (100,000 iterations) for complete data sovereignty.", ACCENT_PURPLE),
        ("Offline QuestPDF Report Card", "Generates professional executive vector PDF reports directly to Documents\\Mirror with zero external service calls.", ACCENT_AMBER)
    ]
    for i, (title, desc, clr) in enumerate(audits):
        row = i // 2; col = i % 2
        x = 0.8 * inch + col * 5.95 * inch
        y = height - 1.8 * inch - (row + 1) * 2.45 * inch
        draw_card(x, y, 5.75 * inch, 2.15 * inch)
        c.setFont("Helvetica-Bold", 15); c.setFillColor(clr); c.drawString(x + 20, y + 2.15 * inch - 32, title)
        c.setFont("Helvetica", 11.5); c.setFillColor(TEXT_MUTED)
        words = desc.split(); lines = []; cur = []
        for w in words:
            if len(" ".join(cur + [w])) < 52: cur.append(w)
            else: lines.append(" ".join(cur)); cur = [w]
        if cur: lines.append(" ".join(cur))
        for li, line in enumerate(lines):
            c.drawString(x + 20, y + 2.15 * inch - 64 - li * 18, line)
    c.showPage()

    # -------------------------------------------------------------
    # SLIDE 10: Winning Proposition
    # -------------------------------------------------------------
    draw_bg()
    draw_card(0.8 * inch, 0.8 * inch, 11.7 * inch, 5.9 * inch, border=ACCENT_CYAN)

    c.setFont("Helvetica-Bold", 11); c.setFillColor(ACCENT_CYAN)
    c.drawString(1.2 * inch, height - 1.4 * inch, "THE QUALCOMM HACKATHON WINNING PROPOSITION")

    c.setFont("Helvetica-Bold", 26); c.setFillColor(TEXT_WHITE)
    c.drawString(1.2 * inch, height - 1.9 * inch, "Why Mirror is the Flagship Showcase for Snapdragon X Copilot+ PCs")

    reasons = [
        ("1. Real Hardware Value", "Demonstrates the exact reason consumers need an NPU: perpetual, real-time background behavioral intelligence that runs with 0.02ms latency and zero battery impact."),
        ("2. Production-Grade Execution", "Not a prototype or concept. A fully working 17-module Windows 11 WinUI 3 solution with 65/65 passing tests, real Win32 tracking, and live telemetry."),
        ("3. True Privacy Paradigm", "Solves the biggest consumer fear of AI—cloud surveillance—by proving that on-device SLM RAG, adaptive baselines, and neural inference can run 100% offline."),
        ("4. Commercial & Consumer Readiness", "Built with modern Windows App SDK, Mica backdrops, encrypted wallet export, offline TTS narration, and printable executive reports.")
    ]
    for idx, (h, d) in enumerate(reasons):
        by = height - 2.6 * inch - idx * 0.95 * inch
        c.setFont("Helvetica-Bold", 12.5); c.setFillColor(ACCENT_EMERALD)
        c.drawString(1.2 * inch, by, f"{h}: ")
        t_w = c.stringWidth(f"{h}: ", "Helvetica-Bold", 12.5)

        c.setFont("Helvetica", 11); c.setFillColor(TEXT_MUTED)
        words = d.split(); lines = []; cur = []
        for w in words:
            if len(" ".join(cur + [w])) < 75: cur.append(w)
            else: lines.append(" ".join(cur)); cur = [w]
        if cur: lines.append(" ".join(cur))
        for li, line in enumerate(lines):
            c.drawString(1.2 * inch + (t_w if li == 0 else 0), by - li * 16, line)

    c.showPage()
    c.save()
    print(f"Saved 10-slide winning PDF with real screenshots to {PDF_PATH}")

if __name__ == "__main__":
    build_winning_pptx()
    build_winning_pdf()
